from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from model import *
OUT="/workspace/dogzie-investisseurs/Dogzie-Universe-projections.xlsx"
wb=Workbook()
BLK=PatternFill("solid",fgColor="0B0905"); DK=PatternFill("solid",fgColor="1A1508"); GLD=PatternFill("solid",fgColor="D4AF37")
INP=PatternFill("solid",fgColor="2A2108")
fT=Font(name="Calibri",size=16,bold=True,color="E2BF57"); fH=Font(bold=True,color="0A0A0A"); fG=Font(bold=True,color="E2BF57")
fW=Font(color="E8E0C8"); fI=Font(color="F3D77A"); fN=Font(italic=True,color="9D937A",size=9); fB=Font(bold=True,color="F1E3B3")
thin=Side(style="thin",color="5A4A1A"); BD=Border(bottom=thin)
CUR='#,##0 "$";-#,##0 "$"'; PCT='0.0%'; INT='#,##0'; PR='#,##0.00 "$"'
def paint(ws,maxr,maxc):
    for r in range(1,maxr+1):
        for c in range(1,maxc+1):
            cell=ws.cell(r,c)
            if cell.fill.fgColor.rgb in (None,"00000000"): cell.fill=BLK
            if cell.font.color is None: cell.font=fW
# ---------- Lisez-moi
ws=wb.active; ws.title="Lisez-moi"
lines=[("DOGZIE UNIVERSE — Projections de revenus (ILLUSTRATIVES)",fT),
("Dogzie Universe (couche blockchain) + Dogzie Town (console PlayStation) + jeton Dogzy. Ancien nom : Family Universe.",fW),("",fW),
("Comment lire ces chiffres",fG),
("• Ce sont des scénarios illustratifs pour discuter avec des investisseurs. Pas une prévision, pas une garantie, pas un audit.",fW),
("• Les contrats intelligents (DogzyToken, PontDogzy) sont rédigés, mais NON déployés et NON audités. Le jeton Dogzy n'est inscrit sur aucune plateforme d'échange.",fW),
("• Toute fonction blockchain exige : déploiement, audit de sécurité indépendant, validation juridique (AMF / ACVM) et validation des règles de Sony PlayStation.",fW),
("• Aucune valeur de jeton ni capitalisation n'est projetée. Les volets crypto sont modélisés en $ CA au moment de la vente (frais / prise de plateforme seulement).",fW),
("• Le capital recherché pour la Phase 2 (500 000 $ à 550 000 $ CA) est un INVESTISSEMENT, pas un revenu. Il n'est inclus dans aucun total.",fW),
("• An 1 = 12 premiers mois après le lancement commercial de Dogzie Town. Devise : $ CA partout.",fW),
("• « Revenu net » = après part PlayStation, commissions de galeries, coût de la marchandise et redevances influenceurs. AVANT salaires, serveurs, développement, marketing payé et impôts. Ce n'est pas un profit.",fW),("",fW),
("Structure du classeur",fG),
("Hypothèses : toutes les hypothèses (cellules dorées = modifiables). Les feuilles de calcul se mettent à jour automatiquement.",fW),
("Prudent / Normal / Optimiste : calcul détaillé par volet et par année (formules).",fW),
("Sommaire : les 9 totaux (3 scénarios x An 1 / Cumul 3 ans / Cumul 5 ans) + annuel An 3 et An 5.",fW),("",fW),
("Repères fournis par la fondatrice (contexte, non vérifiés)",fG),
("143 personnages · 46 bâtiments · 5 extensions · 7 collections jouables · 9 factions · ambiance luxe Dubaï / Saint-Martin",fW),
("Galerie Valkyrie : collection NFT mensuelle + toiles physiques · 23 expositions internationales · 14 curations · Singulart top 10 % (41 000 vues en ~4 mois) · 37 toiles dans le monde (affirmation de la fondatrice)",fW),
("Phase 1 : ~500 $ CA (abonnements) + temps investi (sweat equity) · Phase 2 recherchée : 500 000 $ à 550 000 $ CA",fW),
("Marketing : 5 influenceurs à 1 % de redevances à vie chacun (5 % au total)",fW),
("Économie : accessoires achetés UNIQUEMENT en Dogzy · jetons de mission exportables vers le portefeuille (ex. 500 000 jetons ↔ 5 Dogzy selon la valeur) · export seulement si le personnage existe dans Universe ET Town · autres achats possibles en BNB / autre crypto",fW),
("Volatilité crypto : la valeur en $ CA d'un montant en Dogzy, BNB ou AVAX peut varier fortement entre l'achat et la conversion. Les revenus encaissés en crypto ne sont pas garantis en $ CA.",fN),
]
for i,(t,f) in enumerate(lines,1):
    ws.cell(i,1,t).font=f; ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top")
ws.column_dimensions["A"].width=150; paint(ws,len(lines),3)
# ---------- Hypothèses
wh=wb.create_sheet("Hypothèses")
wh["A1"]="Hypothèses — modifiez les cellules dorées"; wh["A1"].font=fT
for c,h in enumerate(["Hypothèse","Unité","Prudent","Normal","Optimiste","Note"],1):
    x=wh.cell(3,c,h); x.font=fH; x.fill=GLD
ROW={}; r=4; g=None
for k,p in P.items():
    if p["group"]!=g:
        g=p["group"]; wh.cell(r,1,g).font=fG; wh.cell(r,1).fill=DK
        for c in range(2,7): wh.cell(r,c).fill=DK
        r+=1
    wh.cell(r,1,p["label"]); wh.cell(r,2,p["unit"]); wh.cell(r,6,p["note"]).font=fN
    for i in range(3):
        x=wh.cell(r,3+i,p["vals"][i]); x.fill=INP; x.font=fI
        x.number_format=PCT if p["unit"]=="%" else (PR if p["unit"]=="$" and p["vals"][i]<100 and p["vals"][i]%1 else (CUR if p["unit"]=="$" else INT))
    ROW[k]=r; r+=1
r+=1; wh.cell(r,1,"Lieux thématiques actifs par année (croisière An 3, sports An 4, casino-lieu NFT An 5 — aucun jeu d'argent)").font=fG
r+=1; VROW=r
for t in range(5):
    wh.cell(r-1+0,2+t) ; wh.cell(r,2+t,VENUES[t]).fill=INP; wh.cell(r,2+t).font=fI
wh.cell(r,1,"An 1 → An 5 (colonnes B à F)")
r+=2; wh.cell(r,1,"Hors revenus : capital Phase 2 recherché (investissement)").font=fG
wh.cell(r,3,500000).number_format=CUR; wh.cell(r,4,550000).number_format=CUR; wh.cell(r,2,"$ min / max")
wh.column_dimensions["A"].width=72; wh.column_dimensions["B"].width=12
for c in "CDE": wh.column_dimensions[c].width=14
wh.column_dimensions["F"].width=60; wh.freeze_panes="C4"; paint(wh,r,6)
# ---------- scenario sheets
YC=[L(3+t) for t in range(5)]  # C..G
def H(k,s): return f"Hypothèses!${'CDE'[s]}${ROW[k]}"
CALC={}
for s in range(3):
    w=wb.create_sheet(SC[s]); w["A1"]=f"Scénario {SC[s]} — calcul détaillé ($ CA)"; w["A1"].font=fT
    hdr=["Poste","Réf.","An 1","An 2","An 3","An 4","An 5","Cumul 3 ans","Cumul 5 ans"]
    for c,h in enumerate(hdr,1):
        x=w.cell(3,c,h); x.font=fH; x.fill=GLD
    rows=[]; RR={}
    def row(key,label,fn,fmt=CUR,kind="line",cumul=True):
        rows.append((key,label,fn,fmt,kind,cumul))
    row("new","Nouveaux joueurs",lambda t:f"={H(f'new{t+1}',s)}",INT,"vol",False)
    row("active","Joueurs actifs (année précédente x rétention + nouveaux)",lambda t:(f"={YC[t]}{{new}}" if t==0 else f"={YC[t-1]}{{active}}*{H('ret',s)}+{YC[t]}{{new}}"),INT,"vol",False)
    row("hA","A. CONSOLE — DOGZIE TOWN",None,kind="hdr")
    row("A1","Ventes du jeu",lambda t:f"={YC[t]}{{new}}*{H('p_game',s)}")
    row("A2","Extensions (DLC)",lambda t:f"={YC[t]}{{active}}*{H('c_dlc',s)}*{H('p_dlc',s)}*{H('n_dlc',s)}")
    row("A3","Passe compétitive multijoueur",lambda t:f"={YC[t]}{{active}}*{H('c_pass',s)}*{H('p_pass',s)}")
    row("A4","Accessoires / cosmétiques (payés en Dogzy, valeur $ CA)",lambda t:f"={YC[t]}{{active}}*{H('c_acc',s)}*{H('p_acc',s)}")
    row("A","Total A",lambda t:f"=SUM({YC[t]}{{A1}}:{YC[t]}{{A4}})",kind="tot")
    row("hB","B. BLOCKCHAIN — DOGZIE UNIVERSE",None,kind="hdr")
    row("B1m","  NFT personnages vendus (nb)",lambda t:f"={YC[t]}{{active}}*{H('c_nft',s)}+{H('nft_native',s)}",INT,"vol",False)
    row("hold","  Détenteurs cumulés (nb)",lambda t:(f"={YC[t]}{{B1m}}" if t==0 else f"={YC[t-1]}{{hold}}+{YC[t]}{{B1m}}"),INT,"vol",False)
    row("B1","NFT personnages",lambda t:f"={YC[t]}{{B1m}}*{H('p_nft',s)}")
    row("B2","NFT bâtiments / carte",lambda t:f"={H('bldg_lots',s)}*{H('p_bldg',s)}")
    row("B3e","  Éditions Valkyrie par collection (nb)",lambda t:f"=ROUND({H('valk_ed',s)}*(1+{H('valk_g',s)})^{t},0)",INT,"vol",False)
    row("B3","Collections NFT mensuelles Valkyrie (12 / an)",lambda t:f"=12*{YC[t]}{{B3e}}*{H('p_valk',s)}")
    row("B4","Redevances sur reventes secondaires",lambda t:f"=(SUM($C{{B1}}:{YC[t]}{{B1}})+SUM($C{{B2}}:{YC[t]}{{B2}})+SUM($C{{B3}}:{YC[t]}{{B3}}))*{H('turnover',s)}*{H('royalty',s)}")
    row("B5","Frais de plateforme (export de jetons + revente d'accessoires)",lambda t:f"={YC[t]}{{hold}}*{H('c_export',s)}*{H('v_export',s)}*{H('f_export',s)}+{YC[t]}{{A4}}*{H('resale',s)}*{H('f_resale',s)}")
    row("B","Total B",lambda t:f"={YC[t]}{{B1}}+{YC[t]}{{B2}}+{YC[t]}{{B3}}+{YC[t]}{{B4}}+{YC[t]}{{B5}}",kind="tot")
    row("hC","C. ART PHYSIQUE — GALERIE VALKYRIE",None,kind="hdr")
    row("Cn","  Toiles vendues (nb)",lambda t:f"=ROUND({H('canvas',s)}*(1+{H('canvas_g',s)})^{t},0)",INT,"vol",False)
    row("C","Total C — Toiles physiques",lambda t:f"={YC[t]}{{Cn}}*{H('p_canvas',s)}",kind="tot")
    row("hD","D. MARCHANDISE / MARQUE",None,kind="hdr")
    row("D","Total D — Marchandise",lambda t:f"={YC[t]}{{active}}*{H('c_merch',s)}*{H('p_merch',s)}",kind="tot")
    row("hE","E. LIEUX THÉMATIQUES (croisière, sports, casino-lieu NFT sans jeu d'argent)",None,kind="hdr")
    row("E","Total E — Passes NFT de lieux + commandites",lambda t:f"=Hypothèses!{L(2+t)}${VROW}*({YC[t]}{{active}}*{H('c_venue',s)}*{H('p_venue',s)}+{H('sponsor',s)})",kind="tot")
    row("hT","TOTAUX",None,kind="hdr")
    row("gross","REVENU BRUT TOTAL (A+B+C+D+E)",lambda t:f"={YC[t]}{{A}}+{YC[t]}{{B}}+{YC[t]}{{C}}+{YC[t]}{{D}}+{YC[t]}{{E}}",kind="big")
    row("d_sony","− Part PlayStation (30 % sur jeu, DLC, passe)",lambda t:f"=-({YC[t]}{{A1}}+{YC[t]}{{A2}}+{YC[t]}{{A3}})*{H('sony',s)}")
    row("d_gal","− Commissions galeries / plateformes (toiles)",lambda t:f"=-{YC[t]}{{C}}*{H('gallery',s)}")
    row("d_merch","− Coût de production marchandise",lambda t:f"=-{YC[t]}{{D}}*{H('merch_cost',s)}")
    row("d_infl","− Redevances influenceurs (5 % du brut)",lambda t:f"=-{YC[t]}{{gross}}*{H('infl',s)}")
    row("net","REVENU NET POUR L'ENTREPRISE (avant dépenses d'exploitation et impôts)",lambda t:f"={YC[t]}{{gross}}+SUM({YC[t]}{{d_sony}}:{YC[t]}{{d_infl}})",kind="big")
    rr=4
    for key,*_ in rows: RR[key]=rr; rr+=1
    fmtmap={k:"{"+k+"}" for k in RR}
    for key,label,fn,fmt,kind,cumul in rows:
        rr=RR[key]; w.cell(rr,1,label); w.cell(rr,2,key if kind!="hdr" else "")
        w.cell(rr,2).font=fN
        if kind=="hdr":
            for c in range(1,10): w.cell(rr,c).fill=DK
            w.cell(rr,1).font=fG; continue
        for t in range(5):
            f=fn(t)
            for k2,r2 in RR.items(): f=f.replace("{"+k2+"}",str(r2))
            x=w.cell(rr,3+t,f); x.number_format=fmt
        if cumul:
            w.cell(rr,8,f"=SUM(C{rr}:E{rr})").number_format=fmt
            w.cell(rr,9,f"=SUM(C{rr}:G{rr})").number_format=fmt
        if kind=="vol":
            for c in range(1,10): w.cell(rr,c).font=fN
        if kind in("tot","big"):
            for c in range(1,10): w.cell(rr,c).font=fB; w.cell(rr,c).border=BD
        if kind=="big":
            for c in range(1,10): w.cell(rr,c).fill=INP; w.cell(rr,c).font=Font(bold=True,color="F3D77A")
    rr=RR["net"]+2
    w.cell(rr,1,"Le capital Phase 2 (500 000 $ – 550 000 $ CA) est un investissement : il n'apparaît dans aucune ligne de revenu.").font=fN
    w.cell(rr+1,1,"Illustratif, non audité, aucune garantie. Contrats rédigés, non déployés, non audités; jeton non inscrit.").font=fN
    w.column_dimensions["A"].width=70; w.column_dimensions["B"].width=8
    for c in "CDEFGHI": w.column_dimensions[c].width=15
    w.freeze_panes="C4"; paint(w,rr+1,9); CALC[s]=RR
# ---------- Sommaire
wsum=wb.create_sheet("Sommaire",1)
wsum["A1"]="Sommaire — Dogzie Universe (illustratif, $ CA)"; wsum["A1"].font=fT
wsum["A2"]="Scénarios : Prudent · Normal · « Si le concept pogne vraiment » (Optimiste)"; wsum["A2"].font=fN
def ref(s,key,col): return f"='{SC[s]}'!{col}{CALC[s][key]}"
blocks=[("Revenu BRUT total","gross"),("Revenu NET pour l'entreprise (avant dépenses d'exploitation)","net")]
r=4
for title,key in blocks:
    wsum.cell(r,1,title).font=fG; r+=1
    for c,h in enumerate(["Scénario","An 1","Cumul 3 ans","Cumul 5 ans","An 3 (annuel)","An 5 (annuel)"],1):
        x=wsum.cell(r,c,h); x.font=fH; x.fill=GLD
    r+=1
    for s in range(3):
        wsum.cell(r,1,SC[s] if s<2 else "Optimiste (« si ça pogne vraiment »)").font=fB
        for c,col in enumerate(["C","H","I","E","G"],2):
            x=wsum.cell(r,c,ref(s,key,col)); x.number_format=CUR; x.font=fI
        r+=1
    r+=1
wsum.cell(r,1,"Revenu brut cumulatif 5 ans par volet").font=fG; r+=1
for c,h in enumerate(["Volet","Prudent","Normal","Optimiste"],1):
    x=wsum.cell(r,c,h); x.font=fH; x.fill=GLD
r+=1
for k,lab in STREAMS:
    wsum.cell(r,1,f"{k}. {lab}")
    for s in range(3):
        x=wsum.cell(r,2+s,ref(s,k,"I")); x.number_format=CUR
    r+=1
r+=1
wsum.cell(r,1,"Joueurs actifs — An 1 / An 3 / An 5").font=fG; r+=1
for c,h in enumerate(["Scénario","An 1","An 3","An 5"],1):
    x=wsum.cell(r,c,h); x.font=fH; x.fill=GLD
r+=1
for s in range(3):
    wsum.cell(r,1,SC[s])
    for c,col in enumerate(["C","E","G"],2):
        x=wsum.cell(r,c,ref(s,"active",col)); x.number_format=INT
    r+=1
r+=1
wsum.cell(r,1,"Capital Phase 2 recherché : 500 000 $ – 550 000 $ CA = investissement, exclu de tous les totaux.").font=fN; r+=1
wsum.cell(r,1,"ILLUSTRATIF · non audité · pas une garantie · contrats non déployés ni audités · jeton Dogzy non inscrit.").font=fN
wsum.column_dimensions["A"].width=58
for c in "BCDEF": wsum.column_dimensions[c].width=17
paint(wsum,r,6)
for w in wb.worksheets: w.sheet_view.showGridLines=False
wb.save(OUT); print("saved",OUT)
